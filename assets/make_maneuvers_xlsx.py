# STARWROUGHT: scaffold data/maneuvers.xlsx from the roster's Encounter Mode Maneuvers (Mike, T17,
# 2026-10-01: "Let's get rid of all of these Maneuvers, and I will create a Maneuvers.xlsx
# spreadsheet, with 'Enabled?', and other fields, which will now be the authoritative data sheet").
#
# Built for him to edit rather than retype: every row of roster.json's "actions" block becomes a row
# in the actions-workbook shape xlsx_to_trees.py already reads (a _Tree Index of Name | Type | Meta
# note, then one sheet per roster group), the name written plain and its glyphs moved to the Cost
# column, the "Requirements" and "Trigger" lines of the Effect moved to their own columns, and
# Enabled? blank on every row. Blank is the point: by the standing rule a roster row named by a sheet
# is retired under the same document id, so the moment this workbook exists the roster's Maneuvers
# stop shipping to Foundry, and each one returns when Mike writes Yes beside it. The web app and
# the compendium docx keep showing the whole book.
#
# ONE SHOT. Once the file exists it is a SOURCE, Mike's to edit, and this script refuses to touch it
# (--force overwrites, and loses every edit made since; back the file up first). Running it is not a
# pipeline step and build_all.mjs never calls it.
#
#   python assets/make_maneuvers_xlsx.py            -> data/maneuvers.xlsx, if there is none yet
#   python assets/make_maneuvers_xlsx.py --force    -> overwrite it
#   python assets/make_maneuvers_xlsx.py --out X    -> write somewhere else (a dry run into the scratchpad)
import argparse, html, json, os, re, sys
from openpyxl import Workbook
from openpyxl.cell.rich_text import CellRichText, TextBlock
from openpyxl.cell.text import InlineFont
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.styles.colors import Color
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

HERE = os.path.dirname(os.path.abspath(__file__))
ROOT = os.path.dirname(HERE)
DATA = os.path.join(ROOT, "data")
ROSTER = os.path.join(HERE, "roster.json")
ACTIONS_JSON = os.path.join(HERE, "actions.json")
OUT = os.path.join(DATA, "maneuvers.xlsx")
WORKBOOK = os.path.basename(OUT)

# The converter's own name and cost readers, so the two never disagree about what a name is.
sys.path.insert(0, HERE)
from xlsx_to_trees import ACTION_GLYPHS, action_bare, parse_cost  # noqa: E402

# The actions-workbook columns, in the order the converter documents them (first word of each header
# is what it reads). Aura is the 0.5.1 column: "N ft", "N ft allies", "N ft enemies", "visible", "none".
COLUMNS = ["Action", "Cost", "Traits", "Type", "Prerequisites", "Requirements", "Trigger",
           "Description", "Effect", "Automation", "Aura", "Enabled?"]
WIDTHS = {"Action": 18, "Cost": 9, "Traits": 22, "Type": 14, "Prerequisites": 16, "Requirements": 40,
          "Trigger": 30, "Description": 24, "Effect": 90, "Automation": 18, "Aura": 10, "Enabled?": 10}
WRAPPED = ("Requirements", "Trigger", "Description", "Effect")

# Every roster Maneuver is one every character can use, which is what the roster build flags them as
# (build_foundry.mjs `basic: true`), and a sheet action reaches every character's Maneuvers tab when
# its Type begins with "Basic" (`/^basic\b/i`). The price of that one word is grouping: the web app,
# the compendium docx and the Foundry folders group sheet actions by Type, so under this value the
# six roster groups print as one table until the consumers learn to group basic rows by their sheet
# (actions.json keeps the sheet name on every action as `sheet`). The group still names the sheet
# each row sits on and is written to the _Tree Index Meta note. Change TYPE to the group name
# (`TYPE = None`) to keep the headings instead, at the cost of the basic flag.
TYPE = "Basic Action"

# Aura cells the roster text earns on its own: Seek "Covers everything within 30 feet" is centred on
# the seeker and concerns everyone. "Within reach" and "adjacent" get no cell (the reach bands draw
# those), and nothing else in the block reaches N feet around its user.
AURA = {"seek": "30 ft"}

LABEL = re.compile(r"^(Requirements|Trigger)\s+")
TAG = re.compile(r"<(/?)(b|strong|i|em|u|s|strike)\s*>|<span\s+style=\"color:#([0-9a-fA-F]{6})\"\s*>|</span\s*>", re.I)
ANY_TAG = re.compile(r"<[^>]+>")


def split_name(raw):
    """'Strike ❶ to ❸' -> ('Strike', '❶ to ❸'); 'Aid ❶ (⓿↺)' -> ('Aid', '❶ (⓿↺)'); 'Pass ⓿' -> ('Pass', '⓿').
    The plain half must be what the converter would read off the glyphed name, or the roster row of
    that name would not be retired."""
    m = ACTION_GLYPHS.search(raw)
    plain, cost = (raw.strip(), "") if not m else (raw[:m.start()].strip(), raw[m.start():].strip())
    if plain != action_bare(raw):
        sys.exit(f"{raw}: the plain name '{plain}' is not what the converter reads ('{action_bare(raw)}')")
    if cost and parse_cost(cost) is None:
        sys.exit(f"{raw}: the cost '{cost}' does not parse")
    return plain, cost


def split_effect(effect_html):
    """The roster writes a Maneuver's Requirements and Trigger as lines of the Effect ("Requirements
    You have a free hand...<br>Roll Athletics..."). The workbook has columns for both, so a tag-free
    line opening with either label moves there, label off, and one line carrying both ("Trigger You
    fall... Requirements You have a hand free") is split at the second label. Everything else stays
    in the Effect, in order. Returns (requirements, trigger, [effect lines as HTML])."""
    req, trig, lines = [], [], []
    for seg in re.split(r"<br\s*/?>", effect_html):
        s = seg.strip()
        if not s:
            continue
        m = LABEL.match(s)
        if not m or ANY_TAG.search(s):
            lines.append(s)
            continue
        body = s[m.end():].strip()
        other = "Requirements" if m.group(1) == "Trigger" else "Trigger"
        parts = re.split(r"\s+" + other + r"\s+", body, maxsplit=1)
        (trig if m.group(1) == "Trigger" else req).append(parts[0])
        if len(parts) == 2:
            (req if other == "Requirements" else trig).append(parts[1])
    return " ".join(html.unescape(x) for x in req), " ".join(html.unescape(x) for x in trig), lines


def runs_of(lines):
    """Inline HTML (b, i, u, s, span color) over several lines -> [[text, (b, i, u, s, color)], ...].
    Line breaks ride on the run before them and a whitespace-only run is folded into its neighbour,
    because openpyxl marks a run xml:space="preserve" only when it has text of its own, and Excel
    drops an unmarked whitespace run on save: the WELD weld ("PressRequirements") in CLAUDE.md.
    Adjacent runs in the same font are merged, as Excel itself would merge them."""
    runs = []
    state = {"b": False, "i": False, "u": False, "s": False, "color": None}
    key = lambda: (state["b"], state["i"], state["u"], state["s"], state["color"])
    for n, line in enumerate(lines):
        if n:
            if runs: runs[-1][0] += "\n"
            else: runs.append(["\n", key()])
        pos = 0
        for m in TAG.finditer(line):
            text = line[pos:m.start()]
            if text: runs.append([html.unescape(text), key()])
            if m.group(2):
                flag = {"b": "b", "strong": "b", "i": "i", "em": "i", "u": "u", "s": "s", "strike": "s"}[m.group(2).lower()]
                state[flag] = not m.group(1)
            elif m.group(3):
                state["color"] = m.group(3).upper()
            else:
                state["color"] = None
            pos = m.end()
        tail = line[pos:]
        if tail: runs.append([html.unescape(tail), key()])
    stray = [t for t, _ in runs if ANY_TAG.search(t)]
    if stray:
        sys.exit(f"unhandled HTML in a Maneuver effect: {stray[0][:80]}")
    folded = []
    for text, k in runs:
        if not text:
            continue
        if folded and (not text.strip() or folded[-1][1] == k):
            folded[-1][0] += text
        else:
            folded.append([text, k])
    # A leading whitespace-only run, if one ever happened, belongs to the run after it.
    if len(folded) > 1 and not folded[0][0].strip():
        folded[1][0] = folded[0][0] + folded[1][0]; folded.pop(0)
    return folded


PLAIN = (False, False, False, False, None)


def cell_value(lines):
    """A plain string when nothing in the lines is formatted, else a CellRichText."""
    runs = runs_of(lines)
    if all(k == PLAIN for _, k in runs):
        return "".join(t for t, _ in runs)
    blocks = []
    for text, (b, i, u, s, color) in runs:
        if (b, i, u, s, color) == PLAIN:
            blocks.append(text)
        else:
            font = InlineFont(b=b or None, i=i or None, u="single" if u else None, strike=s or None,
                              color=Color(rgb="FF" + color) if color else None)
            blocks.append(TextBlock(font, text))
    return CellRichText(blocks)


def taken_elsewhere():
    """Names another actions workbook already defines (Aid, in data/actions.xlsx), read from the last
    actions.json: the same action in two workbooks is a converter error, so those rows stay the other
    workbook's and are left out here. {bare lower-case name: workbook}."""
    if not os.path.exists(ACTIONS_JSON):
        return {}
    data = json.load(open(ACTIONS_JSON, encoding="utf-8"))
    sources = data.get("source", [])
    sole = sources[0] if len(sources) == 1 else None
    out = {}
    for a in data.get("actions", []):
        wb = a.get("workbook") or sole or "another actions workbook"
        if wb != WORKBOOK:
            out[action_bare(a["name"]).lower()] = wb
    return out


def style_header(ws, columns):
    fill = PatternFill("solid", fgColor="FFEFE6D2")
    for j, name in enumerate(columns, 1):
        c = ws.cell(1, j, name)
        c.font = Font(bold=True)
        c.fill = fill
        c.alignment = Alignment(vertical="top", wrap_text=True)
        ws.column_dimensions[get_column_letter(j)].width = WIDTHS.get(name, 16)
    ws.freeze_panes = "A2"


def build(roster, skip):
    wb = Workbook()
    about = wb.active
    about.title = "About"
    groups = roster.get("actions", {})
    index_rows, sheets, skipped = [], {}, []
    for group, rows in groups.items():
        ws = wb.create_sheet(title=group)
        style_header(ws, COLUMNS)
        enabled_col = get_column_letter(COLUMNS.index("Enabled?") + 1)
        dv = DataValidation(type="list", formula1='"Yes,No"', allow_blank=True, showDropDown=False)
        ws.add_data_validation(dv)
        n = 0
        for raw, traits, effect in rows:
            plain, cost = split_name(raw)
            if plain.lower() in skip:
                skipped.append((plain, skip[plain.lower()])); continue
            req, trig, lines = split_effect(effect)
            if not lines:
                sys.exit(f"{plain}: nothing left for the Effect once its labels moved")
            n += 1
            r = n + 1
            values = {
                "Action": plain, "Cost": cost,
                "Traits": "" if traits.strip() in ("", "—", "-") else traits,
                "Type": TYPE or group, "Prerequisites": "", "Requirements": req, "Trigger": trig,
                "Description": "", "Effect": cell_value(lines), "Automation": "",
                "Aura": AURA.get(plain.lower(), ""), "Enabled?": "",
            }
            for j, name in enumerate(COLUMNS, 1):
                c = ws.cell(r, j)
                c.value = values[name] if values[name] != "" else None
                c.alignment = Alignment(vertical="top", wrap_text=name in WRAPPED)
            dv.add(f"{enabled_col}{r}")
            index_rows.append((plain, TYPE or group, group))
        sheets[group] = n
    idx = wb.create_sheet(title="_Tree Index", index=1)
    style_header(idx, ["Name", "Type", "Meta note"])
    idx.column_dimensions["A"].width = 22; idx.column_dimensions["B"].width = 14; idx.column_dimensions["C"].width = 32
    for r, row in enumerate(index_rows, 2):
        for j, v in enumerate(row, 1):
            idx.cell(r, j, v)
    about.column_dimensions["A"].width = 120
    notes = [
        "STARWROUGHT Maneuvers (Player's Handbook v4.10, Chapter 2). The authoritative sheet for the Encounter Mode Maneuvers (Mike, 2026-10-01).",
        "Generated once by assets/make_maneuvers_xlsx.py from the roster's Maneuver tables, for editing rather than retyping. From here on this file is the source; "
        "the roster row of any Maneuver named here is retired, and the generator will not overwrite this file.",
        "One sheet per group. Columns (first word wins, order free): " + " | ".join(COLUMNS) + ".",
        "Cost takes the handbook's glyphs: ❶ ❷ ❸ (up to ❻), ⓿ for free, ↺ beside a cost for a Reaction (❶↺, ⓿↺), a range as ❶ to ❸ or ❶ or ❸, and ❶ (⓿↺) for a Maneuver whose Reaction half has a cost of its own.",
        "Type: a Type beginning with \"Basic\" puts the Maneuver on every character's Maneuvers tab in Foundry, read from the compendium rather than copied. "
        "The group a row came from is its sheet, and the Meta note on the _Tree Index.",
        "Requirements and Trigger were lifted out of the Effect text where the book printed them as lines; the Effect holds the rest, rich text welcome (bold mirrors everywhere).",
        "Aura: \"N ft\", \"N ft allies\", \"N ft enemies\", add \"visible\" to show the ring by default, or \"none\" when the Effect's \"within N feet\" is centred elsewhere.",
        "Enabled?: Yes ships the row to Foundry; blank or anything else keeps it authored but off the table. Every row starts blank on purpose: nothing ships until it reads Yes. "
        "The web app and the compendium docx show every row regardless.",
    ]
    if skipped:
        notes.append("Left out, because another actions workbook already defines them (the same action in two workbooks is a converter error): "
                     + "; ".join(f"{n} ({w})" for n, w in skipped) + ".")
    for r, text in enumerate(notes, 1):
        c = about.cell(r, 1, text)
        c.alignment = Alignment(wrap_text=True, vertical="top")
    return wb, sheets, skipped


def main():
    for s in (sys.stdout, sys.stderr):
        if hasattr(s, "reconfigure"): s.reconfigure(encoding="utf-8", errors="replace")
    ap = argparse.ArgumentParser(description="Scaffold data/maneuvers.xlsx from roster.json's Maneuver tables. One shot.")
    ap.add_argument("--force", action="store_true", help="overwrite an existing workbook (every edit in it is lost)")
    ap.add_argument("--out", default=OUT, help="write here instead of data/maneuvers.xlsx")
    args = ap.parse_args()
    out = os.path.abspath(args.out)
    lock = os.path.join(os.path.dirname(out), "~$" + os.path.basename(out))
    if os.path.exists(lock):
        sys.exit(f"{os.path.basename(out)} is open in Excel ({os.path.basename(lock)} present); close it first")
    if os.path.exists(out) and not args.force:
        sys.exit(f"{os.path.relpath(out, ROOT)} exists and is a source now (Mike's to edit); not overwritten. "
                 f"--force replaces it and loses every edit made since.")
    roster = json.load(open(ROSTER, encoding="utf-8"))
    wb, sheets, skipped = build(roster, taken_elsewhere())
    wb.save(out)
    total = sum(sheets.values())
    print(f"wrote {os.path.relpath(out, ROOT)}: {total} Maneuvers on {len(sheets)} sheets, Enabled? blank on every row")
    for group, n in sheets.items():
        print(f"  {group}: {n}")
    print("  columns: " + " | ".join(COLUMNS))
    for name, where in skipped:
        print(f"  left out: {name} (already defined in {where})")
    print("Next: python assets/xlsx_to_trees.py, then node assets/build_foundry.mjs. Do not run this script again over an edited file.")


if __name__ == "__main__":
    main()
