# STARWROUGHT data pipeline: data/*.xlsx -> assets/trees.json
# Header-driven: column ORDER doesn't matter; column NAMES do (first word wins).
# Recognized tree-sheet columns: Talent | Tier | Root | Requires | Prerequisites | Description | Effect | Feeds
#                                Grants | Choice | Free Talent | Enabled?
# Recognized index columns:      Tree | Category | Feeds | Flare (triggers) | Meta | Skills
#   Ancestry rows may also carry: Vigor (or HP) | Size | Speed | Senses | Summary
#   -> those generate the "ancestries" block of roster.json, so the sheet owns the chassis.
# Rich text in Description/Effect cells (b/i/u/strike/color) becomes HTML and mirrors everywhere.
# Enabled? (Mike, 2026-10-01, ruling 60): "Yes" ships the row to Foundry; see enabled_flag below.
# The converter keeps every row and only flags it; build_foundry.mjs does the filtering.
#
# An ACTIONS workbook (data/actions.xlsx, Mike, 2026-09-26) is recognised by its _Tree Index
# carrying Name | Type | Meta note instead of Tree | Category. Its other sheets hold one action per
# row and write assets/actions.json; see parse_actions_workbook below for the columns.
import json, os, re, sys, glob
from openpyxl import load_workbook

HERE = os.path.dirname(os.path.abspath(__file__))
ROOT = os.path.dirname(HERE)
DATA = os.path.join(ROOT, "data")
OUT = os.path.join(HERE, "trees.json")
BGOUT = os.path.join(HERE, "backgrounds.json")
ROSTER = os.path.join(HERE, "roster.json")
AOUT = os.path.join(HERE, "actions.json")

def parse_bg_sheet(bws, bgs, warnings, fname):
    bm = header_map(bws, {"name": "background", "rarity": "rarity", "desc": "desc", "effect": "effect", "skills": "skills", "lore": "lore",
                          "enabled": "enabled"})
    if "name" not in bm or "effect" not in bm:
        warnings.append(f"{fname}: Backgrounds sheet needs Background and Effect columns"); return
    for r in bws.iter_rows(min_row=2):
        gv = lambda k: (r[bm[k]] if k in bm and len(r) > bm[k] else None)
        bname = plain(gv("name"))
        if not bname: continue
        eff = cell_html(gv("effect"))
        skills = [x.strip() for x in re.split(r",|•", plain(gv("skills"))) if x.strip()]
        entry = {"name": bname, "rarity": plain(gv("rarity")) or "Common",
                 "enabled": enabled_flag(gv("enabled"), "enabled" in bm),
                 "desc": cell_html(gv("desc")), "effect": eff}
        if skills: entry["skills"] = skills
        lore = plain(gv("lore"))
        if lore: entry["lore"] = lore
        if not skills:  # legacy fallback: parse grants from effect text
            entry["grants"] = parse_grants(re.sub(r"<[^>]+>", "", eff))
        bgs.append(entry)

def parse_lang_sheet(lws, langs, warnings, fname):
    lm = header_map(lws, {"name": "language", "rarity": "rarity", "desc": "desc"})
    if "name" not in lm:
        warnings.append(f"{fname}: Languages sheet needs a Language column"); return
    for r in lws.iter_rows(min_row=2):
        gv = lambda k: (r[lm[k]] if k in lm and len(r) > lm[k] else None)
        lname = plain(gv("name"))
        if not lname: continue
        langs.append({"name": lname, "rarity": plain(gv("rarity")) or "Common", "desc": cell_html(gv("desc"))})

def parse_grants(plain_effect):
    """'trained in either A or B, and you gain 1 Talent Point...' -> [{options:[A,B], extra:1}, ...]"""
    groups = []
    for sent in re.split(r"[.;]", plain_effect):
        pending = None
        for cl in re.split(r",? and you ", sent):
            m = re.search(r"[Tt]rained in (?:either )?(.+)", cl)
            if m:
                clean = []
                for o in re.split(r", or |, | or ", m.group(1)):
                    o = re.sub(r"^(the|a|an) ", "", o.strip()).strip()
                    if not o: continue
                    lm = re.match(r"(.+?) Lore$", o)
                    clean.append(f"Lore ({lm.group(1)})" if lm else o)
                if clean:
                    pending = {"options": clean, "extra": 0}
                    groups.append(pending)
            else:
                em = re.search(r"gain (\d+) [Tt]alent [Pp]oint", cl)
                if em and pending: pending["extra"] += int(em.group(1))
    return groups

# "Save" and "Heritage" are the retired v1.8 spellings of Defense and Bloodline. They stay in the
# tuple so older sheets still convert, and are normalised to the current word on the way in.
CATEGORIES = ("Skill", "Defense", "Save", "Weapon", "Combat Style", "Armor", "Calling",
              "Ancestry", "Culture", "Bloodline", "Heritage", "Background")
CATALIAS = {"Save": "Defense", "Heritage": "Bloodline"}
IDENTITY = ("Ancestry", "Culture", "Bloodline", "Heritage", "Background")  # roots granted by the chargen choice itself
FALLBACK_NEEDED = ("Skill", "Defense", "Weapon", "Combat Style", "Armor")         # tree-level feed required (untrained fallback)
ATTRS = ("Might", "Agility", "Wits", "Presence")

def esc(s):
    return s.replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")

def run_to_html(text, font):
    h = esc(text)
    if font is not None:
        color = getattr(font, "color", None)
        if color is not None and getattr(color, "rgb", None) and isinstance(color.rgb, str) and len(color.rgb) == 8 and color.rgb != "FF000000":
            h = f'<span style="color:#{color.rgb[2:]}">{h}</span>'
        if getattr(font, "strike", None): h = f"<s>{h}</s>"
        if getattr(font, "u", None): h = f"<u>{h}</u>"
        if getattr(font, "i", None): h = f"<i>{h}</i>"
        if getattr(font, "b", None): h = f"<b>{h}</b>"
    return h

def cell_html(cell):
    if cell is None: return ""
    v = cell.value
    if v is None: return ""
    try:
        from openpyxl.cell.rich_text import CellRichText, TextBlock
        if isinstance(v, CellRichText):
            parts = []
            for block in v:
                parts.append(run_to_html(str(block.text), block.font) if isinstance(block, TextBlock) else esc(str(block)))
            return "".join(parts).strip().replace("\n", "<br>")
    except ImportError:
        pass
    h = esc(str(v).strip())
    f = cell.font
    if f is not None:
        if f.strike: h = f"<s>{h}</s>"
        if f.u: h = f"<u>{h}</u>"
        if f.i: h = f"<i>{h}</i>"
        if f.b: h = f"<b>{h}</b>"
    return h.replace("\n", "<br>")

def plain(cell):
    if cell is None: return ""
    v = cell.value
    if v is None: return ""
    return re.sub(r"<[^>]+>", "", str(v)).strip()

def header_map(ws, wanted):
    """wanted: {key: first-word-of-header (lowercase)} -> {key: col_index0} using row 1."""
    m = {}
    for j in range(1, ws.max_column + 1):
        h = plain(ws.cell(1, j)).lower()
        for key, word in wanted.items():
            if key not in m and h.startswith(word):
                m[key] = j - 1
    return m

# ---- The Enabled? column (Mike, 2026-10-01; ruling 60) ---------------------------------------
# Recognised by its first word like every other header. A cell reading "Yes" (any case, trimmed)
# enables the row; anything else (blank, "No", "TBD") disables it. A sheet WITHOUT the column is
# wholly enabled: its absence means the sheet has not been curated yet, not that it is all off
# (Lore and Languages today). The converter DROPS NOTHING: every talent, background and action
# is written with `enabled`, and a tree carries its Root's flag, because a constellation exists in
# play when its Root does. Foundry is the play surface and ships only the enabled rows (ruling 61);
# the web app, the compendium docx and the plates are the authoring views of the whole book.
def enabled_flag(cell, has_column):
    return plain(cell).lower() == "yes" if has_column else True

# Requires-matching strips the cost glyphs off a name: the v4.10 ⓿❶❷❸❹❺❻ and ↺ (with a bracketed
# Reaction cost such as "Aid ❶ (⓿↺)"), the v3 ◆ and ◇, and the capstone star.
GLYPH_CLASS = "◆◇↺★⓿❶❷❸❹❺❻"
def norm(s):
    s = re.sub(r"\(\s*[" + GLYPH_CLASS + r"\s]*\)", "", s or "")
    return re.sub(r"(?:[" + GLYPH_CLASS + r"]|\s+(?:to|or)\s*(?=[" + GLYPH_CLASS + r"])|\s)+$", "", s).strip()

# ---- Actions ---------------------------------------------------------------------------------
# Recognised action-sheet columns (first word wins, order free):
#   Action | Cost | Traits | Type | Prerequisites | Requirements | Trigger | Description | Effect |
#   Automation
# Cost accepts the handbook's glyphs or words: "◆", "◆◆", "◆ to ◆◆◆", "◆ or ◆◆◆", "↺", "◇",
# "1", "1 to 3", "reaction", "free". With no Cost column the glyphs in the Action name are read,
# the way talent names are; with neither, the action costs one action and the converter says so.
# A cell reading "None" is blank: the sheet's own way of saying an action has no prerequisite.
COST_WORDS = {"reaction": "reaction", "free": "0", "passive": "passive", "0": "0",
              "1": "1", "2": "2", "3": "3", "4": "4", "5": "5", "6": "6",
              "one": "1", "two": "2", "three": "3", "four": "4", "five": "5", "six": "6"}
ACTION_GLYPHS = re.compile(r"[◆◇↺★⓿❶❷❸❹❺❻]")

def action_bare(s):
    """Glyphs out (a bracketed Reaction cost such as "(⓿↺)" with them), the trailing cost joiner
    (to/or) out, any run of whitespace to one space. The same normalisation build_foundry.mjs
    applies (actionName), so the two agree on every name."""
    s = re.sub(r"\(\s*[◆◇↺★⓿❶❷❸❹❺❻\s]*\)", "", s or "")
    s = re.sub(r"\s+(to|or)\s*$", "", ACTION_GLYPHS.sub("", s))
    return re.sub(r"\s+", " ", s).strip()

action_key = lambda s: action_bare(s).lower()

def blank_none(s):
    """'None', a dash, or nothing at all is an empty cell. Anything else is kept as written."""
    s = (s or "").strip()
    return "" if s.lower() in ("", "none", "n/a", "-", "—") else s

COST_TOKEN = re.compile(r"[⓿❶❷❸❹❺❻]|◆+|◇")
def _token_value(tok):
    if tok in ("⓿", "◇"): return 0
    if tok[0] == "◆": return min(6, len(tok))
    return "❶❷❸❹❺❻".index(tok) + 1

def parse_cost(text):
    """'❶ to ❸', '❶ or ❸', '⓿↺', '❶ (⓿↺)', '1 or 3', 'reaction' -> {cost, costMax, costMode,
    reaction, reactionCost}; None when it does not parse. Costs are 0..6 as strings, "passive" for
    a Talent with no cost at all; ↺ is the Reaction trait riding beside a cost (a bare ↺, the v3
    form, is a free Reaction). The v3 ◆ (one action per diamond) and ◇ (free) still read."""
    t = (text or "").strip()
    if not t: return None
    reaction = "↺" in t
    reaction_cost = ""
    m = re.search(r"\(([^)]*)\)", t)
    if m and COST_TOKEN.search(m.group(1)):
        inner = COST_TOKEN.findall(m.group(1))
        if inner: reaction_cost = str(_token_value(inner[0]))
        t = t.replace(m.group(0), " ")
    tokens = COST_TOKEN.findall(t)
    if tokens:
        lo, hi = str(_token_value(tokens[0])), str(_token_value(tokens[-1]))
        # Only an "or" between two glyphs is a cost joiner; one in the name ("Hold or Release
        # ❶ to ❸") is not.
        mode = "or" if re.search(r"(?:[⓿❶❷❸❹❺❻◇]|◆+)\s*or\s*(?:[⓿❶❷❸❹❺❻◇]|◆+)", t, re.I) else "to"
        return {"cost": lo, "costMax": hi if (len(tokens) > 1 and hi != lo) else "", "costMode": mode,
                "reaction": reaction, "reactionCost": reaction_cost}
    words = re.sub(r"\s*\bactions?\b", "", t.lower()).replace("↺", "").strip()
    if not words and reaction:
        return {"cost": "0", "costMax": "", "costMode": "to", "reaction": True, "reactionCost": ""}
    parts = re.split(r"\s+(to|or)\s+", words)
    if len(parts) == 1:
        c = COST_WORDS.get(parts[0])
        if c == "reaction": return {"cost": "0", "costMax": "", "costMode": "to", "reaction": True, "reactionCost": ""}
        return {"cost": c, "costMax": "", "costMode": "to", "reaction": reaction, "reactionCost": ""} if c else None
    if len(parts) == 3:
        lo, join, hi = COST_WORDS.get(parts[0]), parts[1], COST_WORDS.get(parts[2])
        if lo is None or hi is None or "reaction" in (lo, hi): return None
        return {"cost": lo, "costMax": hi if hi != lo else "", "costMode": join, "reaction": reaction, "reactionCost": ""}
    return None

def parse_actions_workbook(wb, fname, actions, warnings, errors, defined):
    """One action per row, typed by the row's own Type cell, else its _Tree Index row
    (Name | Type | Meta note), else the sheet the row sits on. Appends to `actions`; errors block
    the write like any other. `defined` is shared across workbooks (bare lower-case name -> file),
    so the same action in two files is an error, as the same tree in two files is."""
    idx = wb["_Tree Index"]
    im = header_map(idx, {"name": "name", "type": "type", "meta": "meta"})
    typed = {}
    for row in idx.iter_rows(min_row=2):
        g = lambda k: blank_none(plain(row[im[k]])) if k in im and len(row) > im[k] else ""
        if g("name"): typed[action_key(g("name"))] = {"type": g("type"), "meta": g("meta")}
    seen, no_cost = {}, []
    for ws in wb.worksheets:
        if ws.title == "_Tree Index": continue
        cm = header_map(ws, {"name": "action", "cost": "cost", "traits": "trait", "type": "type",
                             "prereq": "prereq", "req": "requirement", "trigger": "trigger",
                             "desc": "desc", "effect": "effect", "automation": "automation",
                             "enabled": "enabled"})
        if "name" not in cm:
            warnings.append(f"{fname} / {ws.title}: no Action column, sheet ignored"); continue
        if "effect" not in cm:
            errors.append(f"{fname} / {ws.title}: an action sheet needs Action and Effect columns"); continue
        if "cost" not in cm: no_cost.append(ws.title)
        for r in ws.iter_rows(min_row=2):
            cellv = lambda k: (r[cm[k]] if k in cm and len(r) > cm[k] else None)
            # Every plain read blanks "None", so the rule holds in every column, not just four.
            pv = lambda k: blank_none(plain(cellv(k)))
            raw = plain(cellv("name"))
            if not raw: continue
            # Glyphs may ride in the name, as they do in talent names; the name itself is bare.
            name = action_bare(raw)
            key = name.lower()
            if key in seen:
                errors.append(f"{fname} / {ws.title} / {name}: also defined on sheet '{seen[key]}'"); continue
            if key in defined:
                errors.append(f"{fname} / {ws.title} / {name}: also defined in {defined[key]}"); continue
            seen[key] = ws.title; defined[key] = fname
            cost_text = pv("cost") if "cost" in cm else ""
            if cost_text:
                cost = parse_cost(cost_text)
                if cost is None:
                    errors.append(f"{fname} / {ws.title} / {name}: Cost '{cost_text}' not understood "
                                  f"(❶, ❷, ❸, ⓿, ↺, '❶ to ❸', '1 or 3', reaction, free; the v3 ◆ and ◇ also read)"); continue
            else:
                cost = parse_cost(raw) if ACTION_GLYPHS.search(raw) else None
            defaulted = cost is None
            if defaulted: cost = {"cost": "1", "costMax": "", "costMode": "to", "reaction": False, "reactionCost": ""}
            # The rich-text cells are read through their plain text first, so an Effect of "None"
            # is an empty Effect and a Description of "None" ships blank.
            effect = cell_html(cellv("effect")) if pv("effect") else ""
            if not effect:
                errors.append(f"{fname} / {ws.title} / {name}: empty Effect"); continue
            meta = typed.get(key)
            atype = pv("type") or (meta or {}).get("type") or ws.title
            if meta is None:
                warnings.append(f"{fname} / {ws.title} / {name}: not in _Tree Index, typed by "
                                f"{'its Type cell' if pv('type') else 'its sheet name'}")
            traits = [t for t in (blank_none(x) for x in re.split(r",(?![^(]*\))", pv("traits"))) if t]
            actions.append({"name": name, "type": atype, **cost,
                            "enabled": enabled_flag(cellv("enabled"), "enabled" in cm),
                            "traits": traits,
                            "prerequisites": pv("prereq"), "requirements": pv("req"),
                            "trigger": pv("trigger"), "description": cell_html(cellv("desc")) if pv("desc") else "",
                            "effect": effect, "automation": pv("automation"),
                            "meta": (meta or {}).get("meta", ""), "costDefaulted": defaulted, "sheet": ws.title})
    for k in typed:
        if k not in seen: warnings.append(f"{fname}: '{k}' is in _Tree Index but has no row on any sheet")
    if no_cost:
        warnings.append(f"{fname}: no Cost column on {', '.join(no_cost)}; an action with no glyph in its "
                        f"name costs one action ❶ until the column exists")

# ---- Mike's root rule (v0.38) ----------------------------------------------------------------
# Every root (constellation root and heritage root alike) must (1) hang on something that gets
# ROLLED, so the constellation can Flare from the one talent every member owns, and (2) IMPROVE at
# Expert/Master/Legendary. Warnings, not errors: authoring in progress shouldn't block a sync.
ROLLY = re.compile(r"\b(checks?|saves?|attack rolls?|Strikes?|damage|criticals?|crit|hits?|"
                   r"Aid|Escape|Recall Knowledge|Awareness|Seek|Seeking|AC)\b", re.I)
# Naming a rank is the explicit way to scale; adding "the constellation's proficiency/rank" to
# something is the implicit way, and it scales T->E->M->L by definition. Both count.
SCALES = re.compile(r"\b(Expert|Master|Legendary)\b|"
                    r"(constellation'?s?|this)\s+(proficiency|rank)|(proficiency|rank)\s+to\b", re.I)

# A heading welded to the label that should follow it on its own line: a lowercase letter, a cost
# glyph or a closing bracket running straight into a capitalised rules label.
WELD = re.compile(r"[a-z⓿❶❷❸❹❺❻↺)](Requirements|Trigger|Duration|Frequency|Auditory|Stance)\b")

def root_rule_gripes(tree, nname, effect):
    """Both halves of the root rule, checked against the Effect prose."""
    txt = re.sub(r"<[^>]+>", "", effect)
    out = []
    if not ROLLY.search(txt):
        out.append(f"{tree} / {nname}: root has nothing that gets rolled, so the constellation "
                   f"cannot Flare from its root (root rule, part 1)")
    if not SCALES.search(txt):
        out.append(f"{tree} / {nname}: root never improves at Expert/Master/Legendary "
                   f"(root rule, part 2)")
    return out

def main():
    # The diagnostics carry the action glyphs, and a piped stdout on Windows is cp1252, which cannot
    # encode them: the run would die on its first warning. Reconfigure the streams before printing.
    for s in (sys.stdout, sys.stderr):
        if hasattr(s, "reconfigure"): s.reconfigure(encoding="utf-8", errors="replace")
    files = [f for f in sorted(glob.glob(os.path.join(DATA, "*.xlsx"))) if not os.path.basename(f).startswith("~$")]
    if not files: sys.exit(f"no .xlsx files found in {DATA}")
    out, errors, warnings, sources, bgs, langs, chassis = {}, [], [], {}, [], [], []
    actions, action_files, defined_actions = [], [], {}
    for path in files:
        fname = os.path.basename(path)
        # A workbook open in Excel used to be skipped, which silently wrote a trees.json without
        # its trees. Excel saves atomically, so the file on disk is always the last saved version:
        # read it, and say so.
        if os.path.exists(os.path.join(DATA, "~$" + fname)):
            warnings.append(f"{fname}: open in Excel; converting the last saved version")
        wb = load_workbook(path, rich_text=True)
        if "_Tree Index" in wb.sheetnames:
            probe = header_map(wb["_Tree Index"], {"tree": "tree", "name": "name", "type": "type"})
            if "tree" not in probe and "name" in probe:
                # A Tree-less index with a Name column is an actions workbook. Without a Type column
                # it is an error rather than a fall-through, so a stale actions.json can never ship.
                if "type" not in probe:
                    errors.append(f"{fname}: an actions _Tree Index needs Name and Type columns (Name | Type | Meta note)"); continue
                n_before = len(actions)
                parse_actions_workbook(wb, fname, actions, warnings, errors, defined_actions)
                action_files.append(fname)
                print(f"  {fname}: {len(actions) - n_before} actions")
                continue
        if "_Tree Index" not in wb.sheetnames:
            if "Backgrounds" in wb.sheetnames:
                parse_bg_sheet(wb["Backgrounds"], bgs, warnings, fname)
            if "Languages" in wb.sheetnames:
                parse_lang_sheet(wb["Languages"], langs, warnings, fname)
            if "Backgrounds" not in wb.sheetnames and "Languages" not in wb.sheetnames:
                warnings.append(f"{fname}: no '_Tree Index' sheet; file skipped")
            continue
        idx = wb["_Tree Index"]
        # "Vigor" is the v4.10 name for the per-level chassis number; "HP" still reads.
        im = header_map(idx, {"tree": "tree", "cat": "categ", "feeds": "feeds", "sparks": "flare",
                              "meta": "meta", "skills": "skills", "hp": "hp", "vigor": "vigor", "size": "size",
                              "speed": "speed", "senses": "senses", "summary": "summary", "parent": "parent"})
        if "vigor" in im and "hp" not in im: im["hp"] = im["vigor"]
        if "tree" not in im or "cat" not in im:
            warnings.append(f"{fname}: _Tree Index needs at least 'Tree' and 'Category' columns; file skipped"); continue
        indexed, n_before = set(), len(out)
        for row in idx.iter_rows(min_row=2):
            g = lambda k: plain(row[im[k]]) if k in im and len(row) > im[k] else ""
            name = g("tree")
            if not name: continue
            category = g("cat")
            if category not in CATEGORIES:
                errors.append(f"{fname} / {name}: bad category '{category}' (allowed: {', '.join(CATEGORIES)})"); continue
            category = CATALIAS.get(category, category)   # Save -> Defense, Heritage -> Bloodline
            feeds = g("feeds")
            if feeds and feeds not in ATTRS:
                warnings.append(f"{fname} / {name}: tree Feeds '{feeds}' ignored (not an attribute)"); feeds = ""
            if not feeds and category in FALLBACK_NEEDED:
                errors.append(f"{fname} / {name}: {category} trees need a tree-level Feeds (untrained fallback)"); continue
            if name in out:
                errors.append(f"{name}: defined in both {sources[name]} and {fname}"); continue
            if name not in wb.sheetnames:
                if category == "Background" and "Backgrounds" in wb.sheetnames:
                    indexed.add(name)  # lives in the Backgrounds table, handled below
                else:
                    warnings.append(f"{fname} / {name}: in _Tree Index but no worksheet yet; skipped")
                continue
            indexed.add(name); sources[name] = fname
            ws = wb[name]
            cm = header_map(ws, {"name": "talent", "tier": "tier", "root": "root", "req": "requires",
                                 "prereq": "prereq", "desc": "desc", "effect": "effect", "feeds": "feeds", "grants": "grants",
                                 "choice": "choice", "freetalent": "free", "enabled": "enabled"})
            if "name" not in cm or "tier" not in cm or "effect" not in cm:
                errors.append(f"{fname} / {name}: sheet needs Talent, Tier and Effect columns"); continue
            nodes, root_count, root_name = [], 0, None
            for r in ws.iter_rows(min_row=2):
                cellv = lambda k: (r[cm[k]] if k in cm and len(r) > cm[k] else None)
                pv = lambda k: plain(cellv(k))
                nname = pv("name")
                if not nname: continue
                tier = pv("tier").upper()[:1]
                if tier not in "TEML":
                    errors.append(f"{name} / {nname}: tier must be T/E/M/L"); continue
                rootcell = pv("root").lower()
                is_root = rootcell in ("x", "root", "yes", "true", "1")
                is_hroot = rootcell in ("h", "heritage")
                requires = [s.strip() for s in re.split(r",| or ", pv("req")) if s.strip()]
                prereqs = pv("prereq")
                desc = cell_html(cellv("desc"))
                effect = cell_html(cellv("effect"))
                feeds_o = pv("feeds")
                if not effect:
                    errors.append(f"{name} / {nname}: empty Effect"); continue
                # Excel drops a rich-text run whose whole text is a line break when the file that
                # wrote it did not mark the run xml:space="preserve", and the bold heading before it
                # is welded to the label after it ("PressRequirements"). Caught here so the sheet is
                # mended rather than the weld shipped (review 0.4.1).
                weld = WELD.search(re.sub(r"<[^>]+>", "", re.sub(r"<br\s*/?>", "\n", effect)))
                if weld:
                    warnings.append(f"{name} / {nname}: Effect reads '{weld.group(0)}'; the line break before "
                                    f"'{weld.group(1)}' was lost when the workbook was saved. Mend the cell.")
                # Enabled? (ruling 60): the row's own flag, always written; a sheet without the
                # column is wholly enabled. The tree's flag is its Root's, set once the rows are read.
                node = {"name": nname, "tier": tier, "cost": 1,
                        "enabled": enabled_flag(cellv("enabled"), "enabled" in cm), "effect": effect}
                gr = pv("grants")
                if gr:
                    m2 = re.match(r"^(\d+) in (one|any|different) (Skill|Calling|Combat Style|Armor|Save|Weapon|opened|anywhere)$", gr)
                    m3 = re.match(r"^(\d+) open (non-Skill|anywhere)$", gr)
                    if m2: node["grant"] = {"n": int(m2.group(1)), "mode": m2.group(2), "scope": m2.group(3)}
                    elif m3: node["grant"] = {"n": int(m3.group(1)), "mode": "open", "scope": m3.group(2)}
                    else: warnings.append(f"{name} / {nname}: unparsed Grants '{gr}'")
                # A build-time choice the talent's effect demands ("Choose a Weapon Group"). The
                # cell holds what is being chosen; consumers turn that into a list of options.
                ch = pv("choice")
                if ch: node["choice"] = ch
                # A talent handed over free by this one, with no talent point spent. Drilled is
                # the only one in the book so far: it gives you Weapon Familiarity outright.
                ft = pv("freetalent")
                if ft: node["freeTalent"] = ft
                if requires: node["requires"] = requires
                if desc and desc.upper() != "TBD": node["desc"] = desc
                if prereqs and norm(prereqs) != norm(nname): node["prereqs"] = prereqs
                if feeds_o and feeds_o.upper() != "TBD":
                    if feeds_o not in ATTRS: errors.append(f"{name} / {nname}: bad Feeds '{feeds_o}'")
                    elif feeds_o != feeds: node["feeds"] = feeds_o
                elif feeds_o.upper() == "TBD":
                    warnings.append(f"{name} / {nname}: Feeds 'TBD'; the point feeds no attribute until set")
                if is_root:
                    node["root"] = True; root_count += 1; root_name = nname
                elif is_hroot:
                    node["hroot"] = True
                    if category != "Ancestry":
                        warnings.append(f"{name} / {nname}: bloodline root ('h') outside an Ancestry tree")
                if is_root or is_hroot:
                    warnings.extend(root_rule_gripes(name, nname, effect))
                if "★" in nname and tier != "L":
                    errors.append(f"{name} / {nname}: capstones (★) must be tier L")
                nodes.append(node)
            # requires naming the root are KEPT (they drive the star-map's root edges);
            # the engine enforces root-first regardless of whether a talent lists it.
            if root_count != 1:
                errors.append(f"{name}: needs exactly 1 root (has {root_count})"
                              + (" (identity-tree roots are granted by the chargen choice, but still mark one)" if category in IDENTITY else ""))
            # A constellation exists in play when its Root does (ruling 60), so the tree's flag is
            # the root's. An enabled talent under a disabled root is still written, and warned
            # about: Foundry would ship a Talent with no Root to reach it from.
            tree_enabled = any(n["enabled"] for n in nodes if n.get("root"))
            if not tree_enabled:
                warnings.extend(f"{name} / {n['name']}: enabled, but the constellation's root is not; "
                                f"Foundry has no Root to reach it from" for n in nodes if n["enabled"])
            by_norm = {norm(n["name"]): n for n in nodes}
            for n in nodes:
                for rq in n.get("requires", []):
                    if norm(rq) not in by_norm: errors.append(f"{name} / {n['name']}: requires unknown '{rq}'")
                    # An enabled talent that Requires a disabled one can never be bought in Foundry (ruling 60).
                    elif n["enabled"] and not by_norm[norm(rq)]["enabled"]:
                        warnings.append(f"{name} / {n['name']}: enabled, but Requires '{rq}', which is not")
            label = category + (" constellation" if category in ("Skill", "Defense", "Armor", "Ancestry", "Culture", "Bloodline", "Background") else "")
            meta_extra = g("meta")
            cskills = [x.strip() for x in re.split(r",|•", g("skills")) if x.strip()]
            if category == "Ancestry" and ("hp" in im or "size" in im):
                hp = g("hp")
                hers = [[n["name"], n["effect"]] for n in nodes if n.get("hroot")]
                if not hp.isdigit():
                    errors.append(f"{fname} / {name}: Ancestry Vigor must be a whole number (got '{hp}')")
                elif not hers:
                    errors.append(f"{fname} / {name}: Ancestry has chassis columns but no bloodline roots ('h')")
                else:
                    # `vigor` is the v4.10 key; `hp` is kept one release for readers not yet moved.
                    # `enabled` is the tree's (ruling 60), so chargen can offer only the shipped ancestries.
                    chassis.append({"name": name, "enabled": tree_enabled, "vigor": int(hp), "hp": int(hp), "size": g("size") or "Medium",
                                    "speed": g("speed") or "6 ft", "senses": g("senses") or "—",
                                    "tree": name, "blurb": g("summary"), "bloodlines": hers})
                    if not meta_extra:  # derive the display line when the sheet leaves it blank
                        bits = [g("size"), f"Vigor {hp}", g("speed")] + ([g("senses").lower()] if g("senses") not in ("", "—") else [])
                        meta_extra = " • ".join(b for b in bits if b)
            # A parent Constellation (v4.10): every Talent bought here also counts toward the
            # parent's rank. Melee and Ranged are the parents; the Combat Styles name one of them.
            parent = g("parent")
            out[name] = {"category": category, "enabled": tree_enabled, "feeds": feeds, **({"skills": cskills} if cskills else {}),
                         **({"parent": parent} if parent else {}),
                         "meta": f"{label} • {feeds or '—'}" + (f" ({meta_extra})" if meta_extra else ""),
                         "sparks": g("sparks"), "nodes": nodes}
        if "Backgrounds" in wb.sheetnames:
            parse_bg_sheet(wb["Backgrounds"], bgs, warnings, fname); indexed.add("Backgrounds")
        if "Languages" in wb.sheetnames:
            parse_lang_sheet(wb["Languages"], langs, warnings, fname); indexed.add("Languages")
        extra = [sn for sn in wb.sheetnames if sn != "_Tree Index" and sn not in indexed]
        if extra: warnings.append(f"{fname}: sheets not in its _Tree Index (ignored): {', '.join(extra)}")
        print(f"  {fname}: {len(out) - n_before} trees")
    # A Free Talent may live in another constellation (Drilled hands over Weapon Familiarity, which
    # is authored in Weapons), so it can only be checked once every sheet has been read.
    everywhere = {norm(n["name"]): n for tr in out.values() for n in tr["nodes"]}
    for tname, tree in out.items():
        for n in tree["nodes"]:
            ft = n.get("freeTalent")
            if ft and norm(ft) not in everywhere:
                errors.append(f"{tname} / {n['name']}: Free Talent names no talent in the book: '{ft}'")
            # A disabled Free Talent is one Foundry cannot hand over (ruling 60).
            elif ft and n["enabled"] and not everywhere[norm(ft)]["enabled"]:
                warnings.append(f"{tname} / {n['name']}: enabled, but its Free Talent '{ft}' is not; Foundry cannot hand it over")
        # A Parent must be a tree in the book, and a parent has no parent of its own.
        par = tree.get("parent")
        if par:
            if par not in out:
                errors.append(f"{tname}: Parent names no Constellation in the book: '{par}'")
            elif out[par].get("parent"):
                errors.append(f"{tname}: Parent '{par}' has a parent of its own; only one level of inheritance")
            elif par == tname:
                errors.append(f"{tname}: a Constellation cannot be its own Parent")
    # An enabled Background whose Skills name a constellation with no enabled Root hands out a
    # Training point with nowhere to land in Foundry (ruling 60). A Skill that is not a
    # constellation in the book (a Lore) is not checked, as before.
    for bg in bgs:
        for s in (bg.get("skills", []) if bg["enabled"] else []):
            if s in out and not out[s]["enabled"]:
                warnings.append(f"Backgrounds / {bg['name']}: enabled, but its Skill '{s}' is a constellation whose root is not")
    # What Foundry will ship (ruling 61), on one line beside the per-file counts. Every row is
    # still written; this counts the rows flagged Yes.
    print("enabled for Foundry: "
          f"{sum(1 for t in out.values() if t['enabled'])} of {len(out)} constellations, "
          f"{sum(1 for t in out.values() for n in t['nodes'] if n['enabled'])} of {sum(len(t['nodes']) for t in out.values())} talents, "
          f"{sum(1 for b in bgs if b['enabled'])} of {len(bgs)} backgrounds, "
          f"{sum(1 for a in actions if a['enabled'])} of {len(actions)} actions")
    for w in warnings: print("WARNING:", w)
    if errors:
        print("VALIDATION ERRORS:"); [print("  -", e) for e in errors]; sys.exit(1)
    if not out: sys.exit("no trees found")
    # encoding is explicit everywhere: talent names carry ◆ ↺ ◇ ★, and a Windows default of cp1252
    # would crash the write half way through and leave a truncated trees.json behind.
    json.dump(out, open(OUT, "w", encoding="utf-8"), indent=1, ensure_ascii=False)
    print(f"wrote {OUT}: {len(out)} trees, {sum(len(t['nodes']) for t in out.values())} talents")
    if bgs:
        json.dump(bgs, open(BGOUT, "w", encoding="utf-8"), indent=1, ensure_ascii=False)
        print(f"wrote {BGOUT}: {len(bgs)} backgrounds")
    if langs:
        json.dump(langs, open(os.path.join(HERE, "languages.json"), "w", encoding="utf-8"), indent=1, ensure_ascii=False)
        print(f"wrote languages.json: {len(langs)} languages")
    # Written only when an actions workbook was read, so a checkout without one keeps the file it has.
    if action_files:
        json.dump({"source": action_files, "actions": actions}, open(AOUT, "w", encoding="utf-8"), indent=1, ensure_ascii=False)
        defaulted = [a["name"] for a in actions if a["costDefaulted"]]
        print(f"wrote {AOUT}: {len(actions)} actions from {', '.join(action_files)}"
              + (f" ({len(defaulted)} with no Cost given, costed at one action: {', '.join(defaulted)})" if defaulted else ""))
    elif os.path.exists(AOUT):
        # Printed directly: the warnings list was flushed above, so appending to it here would say nothing.
        print(f"WARNING: no actions workbook in data/; {AOUT} left as found. It is generated from data/actions.xlsx, not hand-kept.")
    # Patch ONLY the ancestries block of roster.json; cultures/weapons/conditions/etc. stay hand-kept.
    # Guarded: an empty parse (chassis columns removed from ancestries.xlsx, or no ancestries workbook
    # in data/ at all) leaves the file alone rather than blanking chargen. A workbook open in Excel is
    # no longer a case here: it is read from its last saved copy above.
    if chassis and os.path.exists(ROSTER):
        roster = json.load(open(ROSTER, encoding="utf-8"))
        before = len(roster.get("ancestries", []))
        if len(chassis) < before:
            print(f"WARNING: roster.json had {before} ancestries, the sheets define {len(chassis)}; "
                  f"the sheets win, so confirm nothing was dropped by accident")
        roster["ancestries"] = chassis   # key keeps its original position in the file
        json.dump(roster, open(ROSTER, "w", encoding="utf-8"), indent=1, ensure_ascii=False)
        print(f"wrote {ROSTER}: {len(chassis)} ancestry chassis from the sheet")

if __name__ == "__main__":
    main()
